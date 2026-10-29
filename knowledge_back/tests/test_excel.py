from copy import copy
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from app.document_authoring import document_bytes
from app.excel_authoring import parse_workbook, workbook_bytes, workbook_text
from app.extensions import db
from app.indexer import _extract_text, process_one
from app.models import (
    DocumentChunk,
    DocumentVersion,
    KnowledgeBase,
    Page,
    Role,
    RolePagePermission,
)

documents = __import__("app.documents", fromlist=["documents"])
indexer = __import__("app.indexer", fromlist=["indexer"])


def login(client, username="writer"):
    response = client.post(
        "/api/v1/auth/login", json={"username": username, "password": "Writer-Password-2026!"}
    )
    assert response.status_code == 200
    return {"Authorization": f"Bearer {response.get_json()['data']['access_token']}"}


def sample_xlsx():
    workbook = Workbook()
    first = workbook.active
    first.title = "设备参数"
    first["A1"] = "名称"
    first["B1"] = "数值"
    first["A2"] = "工作温度"
    first["B2"] = -20
    first["C2"] = "=SUM(B2, 40)"
    first["B2"].number_format = "0.00"
    first_font = copy(first["B1"].font)
    first_font.bold = True
    first_font.color = "FFFF0000"
    first["A1"].font = first_font
    first["B1"].fill = PatternFill(fill_type="solid", fgColor="FFFFFF00")
    first.row_dimensions[1].height = 28
    first.column_dimensions["A"].width = 22
    first.merge_cells("A4:B4")
    first["A4"] = "合并内容"
    first.freeze_panes = "B2"
    second = workbook.create_sheet("说明")
    second["A1"] = "请按设备型号查询参数"
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def sample_xls():
    import xlwt

    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("旧版参数")
    sheet.write(0, 0, "名称")
    sheet.write(1, 0, "工作温度")
    sheet.write(1, 1, -20)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def styled_xlsx():
    workbook = load_workbook(BytesIO(sample_xlsx()))
    cell = workbook.active["D1"]
    cell.value = "样式回归"
    cell.font = Font(bold=True, underline="double", strike=True, color="FFFF0000")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True, text_rotation=135)
    cell.border = Border(bottom=Side(style="thin", color="FF0000FF"))
    cell.fill = PatternFill(fill_type="solid", fgColor="FFFFFF00")
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def test_workbook_round_trip_preserves_sheets_values_formulas_and_layout():
    raw = sample_xlsx()
    workbook = parse_workbook("参数.xlsx", raw)
    assert [sheet["name"] for sheet in workbook["sheets"]] == ["设备参数", "说明"]
    assert workbook["sheets"][0]["cells"]["1,2"]["formula"] == "=SUM(B2, 40)"
    assert workbook["sheets"][0]["cells"]["0,0"]["style"]["font"]["bold"] is True
    assert workbook["sheets"][0]["cells"]["0,0"]["style"]["font"]["color"] == "FFFF0000"
    assert workbook["sheets"][0]["cells"]["0,1"]["style"]["fill"]["color"] == "FFFFFF00"
    assert workbook["sheets"][0]["cells"]["1,1"]["style"]["number_format"] == "0.00"
    assert workbook["sheets"][0]["merges"] == [{"start_row": 3, "start_col": 0, "end_row": 3, "end_col": 1}]
    assert workbook["sheets"][0]["freeze_pane"] == "B2"
    assert "A2：工作温度" in workbook_text(workbook)
    restored = load_workbook(BytesIO(workbook_bytes(workbook)), data_only=False)
    assert restored.sheetnames == ["设备参数", "说明"]
    assert restored["设备参数"]["C2"].value == "=SUM(B2, 40)"
    assert restored["设备参数"]["A4"].value == "合并内容"
    assert restored["设备参数"]["A1"].font.bold is True
    assert restored["设备参数"]["A1"].font.color.rgb == "FFFF0000"
    assert restored["设备参数"]["B1"].fill.fgColor.rgb == "FFFFFF00"
    assert restored["设备参数"]["B2"].number_format == "0.00"
    assert restored["设备参数"].row_dimensions[1].height == 28
    assert restored["设备参数"].column_dimensions["A"].width == 22


def test_document_bytes_and_indexer_extract_excel_text():
    raw = sample_xlsx()
    workbook = parse_workbook("参数.xlsx", raw)
    assert document_bytes("xlsx", workbook).startswith(b"PK")
    text = _extract_text("参数.xlsx", raw)
    assert "工作表：设备参数" in text
    assert "C2：=SUM(B2, 40)" in text


def test_xls_import_can_be_previewed_and_saved_as_xlsx():
    workbook = parse_workbook("旧版参数.xls", sample_xls())
    assert workbook["sheets"][0]["name"] == "旧版参数"
    assert workbook["sheets"][0]["cells"]["1,1"]["value"] == -20.0
    restored = load_workbook(BytesIO(workbook_bytes(workbook)), data_only=False)
    assert restored.sheetnames == ["旧版参数"]
    assert restored["旧版参数"]["A2"].value == "工作温度"
    assert restored["旧版参数"]["B2"].value == -20


@pytest.fixture
def excel_setup(app, client, monkeypatch):
    app.config.update(INDEX_WORKER_AUTOSTART=False)
    result = app.test_cli_runner().invoke(
        args=["create-user", "--username", "writer"],
        input="Writer-Password-2026!\nWriter-Password-2026!\n",
    )
    assert result.exit_code == 0
    with app.app_context():
        role = db.session.scalar(db.select(Role).where(Role.code == "user"))
        for path in ("/create-document", "/documents"):
            page = db.session.scalar(db.select(Page).where(Page.route_path == path))
            db.session.add(RolePagePermission(role_id=role.id, page_id=page.id, can_view=True))
        base = db.session.scalar(db.select(KnowledgeBase))
        base_id = str(base.id)
        db.session.commit()
    objects = {}

    def put(key, stream, *, bucket, length, content_type):
        content = stream.read()
        assert len(content) == length
        objects[(bucket, key)] = content

    monkeypatch.setattr(documents, "put_object", put)
    monkeypatch.setattr(documents, "presign_put", lambda key, *, bucket: f"http://storage/{bucket}/{key}")
    monkeypatch.setattr(
        documents,
        "presign_get",
        lambda key, filename, *, bucket, **kwargs: f"http://storage/{bucket}/{key}",
    )
    monkeypatch.setattr(documents, "read_object", lambda key, *, bucket, max_size: objects[(bucket, key)])
    monkeypatch.setattr(documents, "stat_object", lambda key, *, bucket: SimpleNamespace(size=len(objects[(bucket, key)])))
    monkeypatch.setattr(documents, "delete_object", lambda key, *, bucket: objects.pop((bucket, key), None))
    monkeypatch.setattr(indexer, "read_object", lambda key, *, bucket, max_size: objects[(bucket, key)])
    monkeypatch.setattr(indexer, "embed_texts", lambda texts: ([[1.0, 0.0] for _ in texts], "test-embedding"))
    monkeypatch.setattr(documents, "_schedule_index_task", lambda: None)
    return {"base_id": base_id, "objects": objects, "headers": login(client)}


def test_create_preview_edit_and_reindex_excel(app, client, excel_setup):
    raw = sample_xlsx()
    source = parse_workbook("参数.xlsx", raw)
    document_id = uuid4()
    response = client.post(
        "/api/v1/documents/create",
        headers=excel_setup["headers"],
        json={
            "knowledge_base_id": excel_setup["base_id"],
            "title": "设备参数",
            "format": "xlsx",
            "content": source,
            "request_id": str(document_id),
        },
    )
    assert response.status_code == 201, response.get_json()
    data = response.get_json()["data"]
    assert data["version"]["filename"] == "设备参数.xlsx"
    assert process_one(app)
    preview = client.get(
        f"/api/v1/documents/{document_id}/excel-content?version_id={data['version']['id']}",
        headers=excel_setup["headers"],
    )
    assert preview.status_code == 200
    workbook = preview.get_json()["data"]["workbook"]
    workbook["sheets"][0]["cells"]["1,1"]["value"] = 99
    edit_id = uuid4()
    saved = client.post(
        f"/api/v1/documents/{document_id}/excel-content",
        headers=excel_setup["headers"],
        json={
            "base_version_id": data["version"]["id"],
            "workbook": workbook,
            "request_id": str(edit_id),
        },
    )
    assert saved.status_code == 201, saved.get_json()
    new_version = saved.get_json()["data"]["version"]
    assert new_version["version_no"] == 2
    assert new_version["filename"] == "设备参数.xlsx"
    assert process_one(app)
    updated = client.get(
        f"/api/v1/documents/{document_id}/excel-content?version_id={new_version['id']}",
        headers=excel_setup["headers"],
    )
    assert updated.get_json()["data"]["workbook"]["sheets"][0]["cells"]["1,1"]["value"] == 99
    old = client.get(
        f"/api/v1/documents/{document_id}/excel-content?version_id={data['version']['id']}",
        headers=excel_setup["headers"],
    )
    assert old.get_json()["data"]["workbook"]["sheets"][0]["cells"]["1,1"]["value"] == -20
    assert client.post(
        f"/api/v1/documents/{document_id}/excel-content",
        headers=excel_setup["headers"],
        json={"base_version_id": data["version"]["id"], "workbook": workbook, "request_id": str(edit_id)},
    ).status_code == 200
    with app.app_context():
        assert db.session.scalar(db.select(db.func.count()).select_from(DocumentVersion)) == 2
        assert db.session.scalar(db.select(DocumentChunk)) is not None


def test_excel_upload_is_allowlisted_and_corrupt_content_is_rejected(app, client, excel_setup):
    response = client.post(
        "/api/v1/documents/presign-upload",
        headers=excel_setup["headers"],
        json={
            "knowledge_base_id": excel_setup["base_id"],
            "filename": "坏文件.xlsx",
            "file_size": 4,
            "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        },
    )
    assert response.status_code == 201
    version_id = response.get_json()["data"]["version"]["id"]
    upload = client.put(
        f"/api/v1/documents/versions/{version_id}/upload",
        headers={**excel_setup["headers"], "Content-Type": "application/octet-stream"},
        data=b"bad!",
    )
    assert upload.status_code == 422


@pytest.mark.parametrize("sample,extension", [(sample_xlsx, "xlsx"), (sample_xls, "xls")])
def test_excel_upload_completion_keeps_original_bytes_for_download(app, client, excel_setup, sample, extension):
    raw = sample()
    prepared = client.post(
        "/api/v1/documents/presign-upload",
        headers=excel_setup["headers"],
        json={
            "knowledge_base_id": excel_setup["base_id"],
            "filename": f"原始格式.{extension}",
            "title": "原始格式",
            "file_size": len(raw),
            "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        },
    )
    assert prepared.status_code == 201
    ticket = prepared.get_json()["data"]
    uploaded = client.put(
        f"/api/v1/documents/versions/{ticket['version']['id']}/upload",
        headers={**excel_setup["headers"], "Content-Type": "application/octet-stream"},
        data=raw,
    )
    assert uploaded.status_code == 202, uploaded.get_json()
    stored = excel_setup["objects"][(ticket["bucket"], ticket["object_key"])]
    assert stored == raw
    download = client.get(
        f"/api/v1/documents/{ticket['document']['id']}/download-url?version_id={ticket['version']['id']}",
        headers=excel_setup["headers"],
    )
    assert download.status_code == 200
    assert download.get_json()["data"]["download_url"].endswith(ticket["object_key"])


@pytest.mark.parametrize("style", [
    {"alignment": {"vertical": 2}},
    {"alignment": {"horizontal": "invalid"}},
    {"border": {"bottom": {"style": 1}}},
    {"font": {"underline": 1}},
    {"alignment": []},
    {"font": "invalid"},
])
def test_create_excel_rejects_invalid_styles_without_server_error(client, excel_setup, style):
    response = client.post(
        "/api/v1/documents/create",
        headers=excel_setup["headers"],
        json={
            "knowledge_base_id": excel_setup["base_id"], "title": "样式校验", "format": "xlsx",
            "content": {"sheets": [{"name": "Sheet1", "cells": {"0,0": {"value": "内容", "style": style}}}]},
            "request_id": str(uuid4()),
        },
    )
    assert response.status_code == 422, response.get_json()
    assert "样式" in response.get_json()["message"]
    assert excel_setup["objects"] == {}


def test_text_renamed_as_excel_has_actionable_upload_error(client, excel_setup):
    raw = "部门\\月份\t1\t2\n开发部门\t123\t456\n".encode()
    prepared = client.post(
        "/api/v1/documents/presign-upload", headers=excel_setup["headers"],
        json={"knowledge_base_id": excel_setup["base_id"], "filename": "收支表.xlsx", "file_size": len(raw)},
    )
    assert prepared.status_code == 201, prepared.get_json()
    version_id = prepared.get_json()["data"]["version"]["id"]
    uploaded = client.put(
        f"/api/v1/documents/versions/{version_id}/upload",
        headers={**excel_setup["headers"], "Content-Type": "application/octet-stream"}, data=raw,
    )
    assert uploaded.status_code == 422
    assert "实际是文本文件" in uploaded.get_json()["message"]
    assert "另存为" in uploaded.get_json()["message"]
