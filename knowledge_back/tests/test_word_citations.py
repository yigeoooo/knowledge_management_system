import importlib
from io import BytesIO
from uuid import UUID, uuid4

import test_document_creation as creation_tests
from docx import Document as Docx
from docx.shared import RGBColor
from test_document_creation import create, word_content
from test_qa_logs import admin_headers

from app.extensions import db
from app.models import DocumentVersion, KnowledgeBase, QALog
from app.retrieval import with_preview_urls

authoring_setup = creation_tests.creation


def test_word_citations_render_old_version_formats_in_live_and_historical_views(
    app, client, authoring_setup, monkeypatch
):
    response, _ = create(client, authoring_setup, format="docx", content=word_content())
    data = response.get_json()["data"]
    document_id, version_id = data["document"]["id"], data["version"]["id"]
    retrieval = importlib.import_module("app.retrieval")
    monkeypatch.setattr(
        retrieval,
        "read_object",
        lambda key, **kw: authoring_setup["objects"][(kw["bucket"], key)],
        raising=False,
    )
    source = {
        "document_id": document_id,
        "version_id": version_id,
        "filename": "guide.docx",
        "content": "检查电源\n\n启动设备",
        "modality": "text",
    }
    with app.app_context():
        rich = with_preview_urls([source])[0]
        assert "font-weight:700" in rich["content_html"]
        assert "<ol" in rich["content_html"]
        assert "绿色表示就绪" not in rich["content_html"]  # Only the cited range.
        base = db.session.scalar(db.select(KnowledgeBase))
        log = QALog(
            request_id=str(uuid4()),
            source_type="WEB",
            knowledge_base_id=base.id,
            question="检查什么？",
            answer="电源",
            status="SUCCESS",
            sources=[source],
        )
        db.session.add(log)
        db.session.commit()
        log_id = str(log.id)
    historical = client.get(f"/api/v1/qa-logs/{log_id}", headers=admin_headers(client))
    assert historical.status_code == 200
    assert historical.get_json()["data"]["sources"][0]["content_html"] == rich["content_html"]
    with app.app_context():
        assert db.session.get(QALog, UUID(log_id)).sources == [source]


def test_citation_range_does_not_expose_other_parts_of_same_paragraph(
    app, client, authoring_setup, monkeypatch
):
    response, _ = create(client, authoring_setup, format="docx", content=word_content())
    data = response.get_json()["data"]
    doc = Docx()
    p = doc.add_paragraph()
    p.add_run("无关前文")
    r = p.add_run("引用范围")
    r.bold = True
    r.font.color.rgb = RGBColor(255, 0, 0)
    p.add_run("无关后文")
    output = BytesIO()
    doc.save(output)
    retrieval = importlib.import_module("app.retrieval")
    monkeypatch.setattr(retrieval, "read_object", lambda *a, **kw: output.getvalue(), raising=False)
    source = {
        "document_id": data["document"]["id"],
        "version_id": data["version"]["id"],
        "filename": "range.docx",
        "content": "引用范围",
        "modality": "text",
    }
    with app.app_context():
        result = with_preview_urls([source])[0]
        assert "引用范围" in result["content_html"] and "font-weight:700" in result["content_html"]
        assert "无关前文" not in result["content_html"] and "无关后文" not in result["content_html"]
        version = db.session.get(DocumentVersion, UUID(data["version"]["id"]))
        version.document_id = uuid4()
        assert "content_html" not in with_preview_urls([source])[0]
