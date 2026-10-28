import copy
import importlib

import test_document_creation as creation_tests
from io import BytesIO
from uuid import UUID, uuid4

import pytest
from docx import Document as Docx
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor
from test_document_creation import create
from test_document_creation import word_content as authored_content

from app.document_authoring import document_bytes
from app.extensions import db
from app.http import ApiError
from app.indexer import process_one
from app.models import Document, DocumentVersion, IndexTask, Page, Role, RolePagePermission
from app.storage import StorageError
from app.word_import import word_content

routes = importlib.import_module('app.documents')

authoring_setup = creation_tests.creation


def external_word():
    doc = Docx()
    doc.sections[0].orientation = WD_ORIENT.LANDSCAPE
    doc.sections[0].page_width, doc.sections[0].page_height = Cm(29.7), Cm(21)
    doc.sections[0].header.paragraphs[0].text = '保留原文件页眉'
    doc.add_heading('设备维护', level=1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(21)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run('红色加粗斜体')
    run.bold, run.italic, run.underline = True, True, True
    run.font.name, run.font.size, run.font.color.rgb = 'SimSun', Pt(18), RGBColor(255, 0, 0)
    doc.add_paragraph('第一项', style='List Number')
    doc.add_paragraph('第二项', style='List Number')
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(1, 0)).text = '跨行'
    table.cell(0, 1).merge(table.cell(0, 2)).text = '跨列'
    table.cell(1, 1).text = '普通单元格'
    table.cell(1, 2).text = '第二列'
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph('第二页')
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def test_external_docx_import_edit_export_preserves_formats_and_original():
    original = external_word()
    content = word_content(original)
    assert content['attrs']['page']['orientation'] == 'landscape'
    assert content['content'][0]['type'] == 'heading'
    p = content['content'][1]
    assert p['attrs']['textAlign'] == 'center'
    assert p['attrs']['lineHeight'] == 1.5
    p['content'][0]['text'] = '修改后的红色标题'
    result = Docx(BytesIO(document_bytes('docx', content, source=original)))
    run = result.paragraphs[1].runs[0]
    assert run.bold and run.italic and run.underline
    assert str(run.font.color.rgb) == 'FF0000' and run.font.size.pt == 18
    assert result.paragraphs[1].paragraph_format.first_line_indent.pt == 21
    assert result.paragraphs[1].paragraph_format.line_spacing == 1.5
    assert result.sections[0].header.paragraphs[0].text == '保留原文件页眉'
    assert result.tables[0].cell(1, 0).text == '跨行'
    assert result.tables[0].cell(0, 2).text == '跨列'
    assert result.element.xpath('.//w:br[@w:type="page"]')
    assert Docx(BytesIO(original)).paragraphs[1].text == '红色加粗斜体'


def test_editor_snapshot_round_trip_and_external_edit_invalidation():
    content = authored_content()
    payload = document_bytes('docx', content)
    assert word_content(payload) == content
    doc = Docx(BytesIO(payload))
    doc.paragraphs[0].runs[0].text = '在桌面 Word 中修改过'
    output = BytesIO()
    doc.save(output)
    imported = word_content(output.getvalue())
    assert imported['content'][0]['content'][0]['text'] == '在桌面 Word 中修改过'


def test_unsupported_content_is_explicit_instead_of_silently_discarded():
    doc = Docx()
    doc.add_paragraph('正文').add_run()._r.append(OxmlElement('w:drawing'))
    output = BytesIO()
    doc.save(output)
    with pytest.raises(ApiError) as error:
        word_content(output.getvalue())
    assert '图片' in error.value.message
    with pytest.raises(ApiError):
        word_content(b'broken docx')


def grant_documents(app):
    with app.app_context():
        role = db.session.scalar(db.select(Role).where(Role.code == 'user'))
        page = db.session.scalar(db.select(Page).where(Page.route_path == '/documents'))
        db.session.add(RolePagePermission(role_id=role.id, page_id=page.id, can_view=True))
        db.session.commit()


def test_word_create_edit_preview_download_index_and_retry(app, client, authoring_setup, monkeypatch):
    grant_documents(app)
    response, _ = create(client, authoring_setup, format='docx', content=authored_content())
    saved = response.get_json()['data']
    document_id, base_id = saved['document']['id'], saved['version']['id']
    original_objects = copy.deepcopy(authoring_setup['objects'])
    path = f'/api/v1/documents/{document_id}/word-content'
    headers = authoring_setup['headers']
    opened = client.get(f'{path}?version_id={base_id}', headers=headers)
    assert opened.status_code == 200
    content = opened.get_json()['data']['content']
    content['content'][1]['content'][0]['text'] = '修改后检查电源'
    body = {'base_version_id': base_id, 'content': content, 'request_id': str(uuid4())}
    edited = client.post(path, headers=headers, json=body)
    assert edited.status_code == 201, edited.get_json()
    version_id = edited.get_json()['data']['version']['id']
    assert edited.get_json()['data']['version']['version_no'] == 2
    assert client.post(path, headers=headers, json=body).status_code == 200
    stale = client.post(path, headers=headers, json={**body, 'request_id': str(uuid4())})
    assert stale.status_code == 409
    assert all(authoring_setup['objects'][key] == value for key, value in original_objects.items())
    assert len(authoring_setup['objects']) == 2
    with app.app_context():
        version = db.session.get(DocumentVersion, UUID(version_id))
        payload = authoring_setup['objects'][(version.minio_bucket, version.minio_object_key)]
        assert Docx(BytesIO(payload)).paragraphs[1].runs[0].bold
        assert Docx(BytesIO(payload)).paragraphs[1].text == '修改后检查电源'
        key = version.minio_object_key
        assert db.session.scalar(db.select(db.func.count()).select_from(IndexTask)) == 2
    monkeypatch.setattr(routes, 'presign_get', lambda object_key, *args, **kwargs: f'https://storage.test/{object_key}')
    download = client.get(f'/api/v1/documents/{document_id}/download-url?version_id={version_id}', headers=headers)
    assert download.get_json()['data']['download_url'].endswith(key)
    preview = client.get(f'/api/v1/documents/{document_id}/text-preview?version_id={version_id}', headers=headers)
    assert preview.status_code == 200 and '修改后检查电源' in preview.get_json()['data']['html']
    reopened = client.get(f'{path}?version_id={version_id}', headers=headers)
    assert reopened.get_json()['data']['content'] == content
    assert process_one(app) and process_one(app)
    with app.app_context():
        assert db.session.get(Document, UUID(document_id)).current_version_id == UUID(version_id)


def test_word_save_failure_preserves_source_and_permissions(app, client, authoring_setup, monkeypatch):
    response, _ = create(client, authoring_setup, format='docx', content=authored_content())
    data = response.get_json()['data']
    path = f"/api/v1/documents/{data['document']['id']}/word-content"
    body = {'base_version_id': data['version']['id'], 'content': authored_content(), 'request_id': str(uuid4())}
    assert client.get(f"{path}?version_id={body['base_version_id']}", headers=authoring_setup['headers']).status_code == 403
    assert client.post(path, headers=authoring_setup['headers'], json=body).status_code == 403
    grant_documents(app)
    monkeypatch.setattr(routes, 'stat_object', lambda *a, **kw: (_ for _ in ()).throw(StorageError('存储不可用')))
    assert client.post(path, headers=authoring_setup['headers'], json=body).status_code == 503
    assert len(authoring_setup['objects']) == 1
    with app.app_context():
        assert db.session.scalar(db.select(db.func.count()).select_from(DocumentVersion)) == 1


def test_200_mb_upload_limit_and_flask_body_limit(app, client, authoring_setup, monkeypatch):
    from app import config

    monkeypatch.setattr(config, 'load_dotenv', lambda *a: None)
    monkeypatch.delenv('MAX_DOCUMENT_SIZE', raising=False)
    defaults = config.configuration({'JWT_SECRET_KEY': 'size-test-secret-32-characters-long'})
    limit = 200 * 1024 * 1024
    assert defaults['MAX_DOCUMENT_SIZE'] == defaults['MAX_CONTENT_LENGTH'] == limit
    monkeypatch.setenv('MAX_DOCUMENT_SIZE', '314572800')
    custom = config.configuration({'JWT_SECRET_KEY': 'size-test-secret-32-characters-long'})
    assert custom['MAX_DOCUMENT_SIZE'] == custom['MAX_CONTENT_LENGTH'] == 300 * 1024 * 1024
    grant_documents(app)
    app.config.update(MAX_DOCUMENT_SIZE=limit, MAX_CONTENT_LENGTH=limit)
    monkeypatch.setattr(routes, 'presign_put', lambda *a, **kw: 'https://storage.test/upload')
    body = {'knowledge_base_id': authoring_setup['base_id'], 'filename': 'large.docx', 'file_size': limit}
    response = client.post('/api/v1/documents/presign-upload', headers=authoring_setup['headers'], json=body)
    assert response.status_code == 201
    version_id = response.get_json()['data']['version']['id']
    too_large = client.post('/api/v1/documents/presign-upload', headers=authoring_setup['headers'], json={**body, 'file_size': limit + 1})
    assert too_large.status_code == 422
    listed = client.get(f"/api/v1/knowledge-bases/{authoring_setup['base_id']}/documents", headers=authoring_setup['headers'])
    assert listed.get_json()['data']['max_file_size'] == limit
    # Exercise Flask's body-size guard without allocating a 200 MB test object.
    oversized = client.open(f'/api/v1/documents/versions/{version_id}/upload', method='PUT',
        headers=authoring_setup['headers'], input_stream=BytesIO(b'x'),
        environ_overrides={'CONTENT_LENGTH': str(limit + 1)})
    assert oversized.status_code in (413, 422)
