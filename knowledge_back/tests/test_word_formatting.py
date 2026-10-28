from io import BytesIO

import pytest
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from app.document_authoring import document_bytes
from app.docx_preview import docx_html
from app.http import ApiError


def paragraph(text, attrs=None, marks=None):
    return {
        "type": "paragraph",
        "attrs": attrs or {},
        "content": [{"type": "text", "text": text, "marks": marks or []}],
    }


def document(*blocks, page=None):
    return {"type": "doc", "content": list(blocks), "attrs": {"page": page} if page else {}}


def read(content):
    return Document(BytesIO(document_bytes("docx", content)))


def test_umo_alignment_typography_and_spacing_survive_docx_export():
    result = read(
        document(
            paragraph(
                "居中标题",
                {
                    "textAlign": "center",
                    "lineHeight": 1.5,
                    "indent": 1,
                    "indentUnit": "em",
                    "margin": {"top": "8", "bottom": "16"},
                },
                [
                    {
                        "type": "textStyle",
                        "attrs": {
                            "fontFamily": "SimSun",
                            "fontSize": "18pt",
                            "color": "rgb(18, 52, 86)",
                            "backgroundColor": "#ffff8a",
                        },
                    },
                    {"type": "letterSpacing", "attrs": {"spacing": "2px"}},
                    {"type": "bold"},
                ],
            )
        )
    )
    p = result.paragraphs[0]
    assert p.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert p.paragraph_format.line_spacing == 1.5
    assert p.paragraph_format.first_line_indent.pt == 24
    assert p.paragraph_format.space_before.pt == 6
    assert p.paragraph_format.space_after.pt == 12
    run = p.runs[0]
    assert run.font.name == "SimSun" and run.font.size.pt == 18 and run.bold
    assert run._r.rPr.rFonts.get(qn("w:eastAsia")) == "SimSun"
    assert str(run.font.color.rgb) == "123456"
    assert run._r.rPr.find(qn("w:shd")).get(qn("w:fill")) == "FFFF8A"
    assert run._r.rPr.find(qn("w:spacing")).get(qn("w:val")) == "30"


def test_merged_tables_preserve_text_background_and_alignment():
    def cell(text, **attrs):
        return {"type": "tableCell", "attrs": attrs, "content": [paragraph(text)]}

    content = document(
        {
            "type": "table",
            "content": [
                {
                    "type": "tableRow",
                    "content": [
                        cell(
                            "跨行",
                            rowspan=2,
                            align="center-middle",
                            background="#edf",
                            colwidth=[100],
                        ),
                        cell("跨列", colspan=2),
                    ],
                },
                {"type": "tableRow", "content": [cell("甲"), cell("乙")]},
            ],
        }
    )
    table = read(content).tables[0]
    assert len(table.columns) == 3 and len(table.rows) == 2
    assert table.cell(0, 0).text == "跨行" and table.cell(1, 0).text == "跨行"
    assert table.cell(0, 1).text == "跨列" and table.cell(0, 2).text == "跨列"
    assert table.cell(1, 1).text == "甲" and table.cell(1, 2).text == "乙"
    assert table.cell(0, 0).vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.CENTER
    assert table.cell(0, 0).paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert table.cell(0, 0)._tc.tcPr.find(qn("w:shd")).get(qn("w:fill")) == "EEDDFF"


def test_page_settings_breaks_tasks_and_superscript_are_saved():
    result = read(
        document(
            paragraph("平方", marks=[{"type": "superscript"}]),
            {"type": "pageBreak"},
            {
                "type": "taskList",
                "content": [
                    {
                        "type": "taskItem",
                        "attrs": {"checked": True},
                        "content": [paragraph("完成检查")],
                    }
                ],
            },
            page={
                "size": {"width": 21, "height": 29.7},
                "orientation": "landscape",
                "margin": {"left": 2, "right": 2, "top": 2.54, "bottom": 2.54},
            },
        )
    )
    section = result.sections[0]
    assert section.orientation == WD_ORIENT.LANDSCAPE
    assert section.page_width.cm == pytest.approx(29.7, abs=0.01)
    assert section.left_margin.cm == pytest.approx(2, abs=0.01)
    assert result.paragraphs[0].runs[0].font.superscript
    assert result.paragraphs[1]._p.xpath('.//w:br[@w:type="page"]')
    assert result.paragraphs[2].text == "☑ 完成检查"


def test_docx_preview_preserves_rich_formatting_and_escapes_text():
    content = document(
        {"type": "heading", "attrs": {"level": 1}, "content": [{"type": "text", "text": "标题"}]},
        paragraph(
            "重要 <script>alert(1)</script>",
            {"textAlign": "center", "lineHeight": 1.5},
            [{"type": "bold"}, {"type": "textStyle", "attrs": {"color": "#ff0000"}}],
        ),
        {
            "type": "orderedList",
            "content": [{"type": "listItem", "content": [paragraph("列表项")]}],
        },
        {
            "type": "table",
            "content": [{"type": "tableRow", "content": [{"type": "tableCell", "content": [paragraph("单元格")]}]}],
        },
    )
    html = docx_html(document_bytes("docx", content))
    assert "<h1" in html and "标题" in html
    assert "font-weight:700" in html and "color:#FF0000" in html
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in html
    assert "<ol " in html and "列表项" in html
    assert '<table class="docx-preview-table">' in html and "单元格" in html


@pytest.mark.parametrize(
    "attrs",
    [
        {"textAlign": []},
        {"textAlign": "unsafe"},
        {"lineHeight": "NaN"},
        {"lineHeight": {}},
        {"indent": -1},
        {"indentUnit": [], "indent": 1},
        {"margin": []},
        {"fontSize": "url(file:///etc/passwd)"},
        {"color": "rgb(999, 0, 0)"},
        {"fontFamily": "bad\x00font"},
    ],
)
def test_invalid_formatting_returns_validation_error(attrs):
    with pytest.raises(ApiError) as error:
        document_bytes("docx", document(paragraph("正文", attrs)))
    assert error.value.status == 422


def test_invalid_table_span_is_rejected():
    with pytest.raises(ApiError):
        read(
            document(
                {
                    "type": "table",
                    "content": [
                        {
                            "type": "tableRow",
                            "content": [
                                {
                                    "type": "tableCell",
                                    "attrs": {"rowspan": 2},
                                    "content": [paragraph("超出范围")],
                                },
                            ],
                        }
                    ],
                }
            )
        )
