import importlib
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import Mock
from uuid import UUID, uuid4

import pytest
from docx import Document as WordDocument

from app.extensions import db
from app.indexer import _extract_text, process_one
from app.models import (
    AuditLog,
    Document,
    DocumentChunk,
    DocumentVersion,
    IndexTask,
    KnowledgeBase,
    Page,
    Role,
    RolePagePermission,
)
from app.storage import StorageError

documents = importlib.import_module("app.documents")
indexer = importlib.import_module("app.indexer")


def word_content():
    def paragraph(text):
        return {"type": "paragraph", "content": [{"type": "text", "text": text}]}

    return {
        "type": "doc",
        "content": [
            {
                "type": "heading",
                "attrs": {"level": 1},
                "content": [{"type": "text", "text": "设备指南"}],
            },
            {
                "type": "paragraph",
                "content": [
                    {
                        "type": "text",
                        "text": "检查电源",
                        "marks": [{"type": "bold"}, {"type": "underline"}],
                    }
                ],
            },
            {
                "type": "orderedList",
                "attrs": {"start": 3},
                "content": [{"type": "listItem", "content": [paragraph("启动设备")]}],
            },
            {
                "type": "table",
                "content": [
                    {
                        "type": "tableRow",
                        "content": [
                            {"type": "tableHeader", "content": [paragraph("项目")]},
                            {"type": "tableHeader", "content": [paragraph("说明")]},
                        ],
                    },
                    {
                        "type": "tableRow",
                        "content": [
                            {"type": "tableCell", "content": [paragraph("指示灯")]},
                            {"type": "tableCell", "content": [paragraph("绿色表示就绪")]},
                        ],
                    },
                ],
            },
        ],
    }


def login(client, username="writer"):
    response = client.post(
        "/api/v1/auth/login", json={"username": username, "password": "Writer-Password-2026!"}
    )
    assert response.status_code == 200
    return {"Authorization": f"Bearer {response.get_json()['data']['access_token']}"}


@pytest.fixture
def creation(app, client, monkeypatch):
    app.config.update(INDEX_WORKER_AUTOSTART=False)
    result = app.test_cli_runner().invoke(
        args=["create-user", "--username", "writer"],
        input="Writer-Password-2026!\nWriter-Password-2026!\n",
    )
    assert result.exit_code == 0
    with app.app_context():
        role = db.session.scalar(db.select(Role).where(Role.code == "user"))
        page = db.session.scalar(db.select(Page).where(Page.route_path == "/create-document"))
        db.session.add(RolePagePermission(role_id=role.id, page_id=page.id, can_view=True))
        base = db.session.scalar(db.select(KnowledgeBase))
        base_id = str(base.id)
        db.session.commit()
    objects = {}

    def put(key, stream, *, bucket, length, content_type):
        content = stream.read()
        assert len(content) == length
        objects[(bucket, key)] = content

    monkeypatch.setattr(documents, "put_object", put)
    monkeypatch.setattr(
        documents, "read_object", lambda key, *, bucket, max_size: objects[(bucket, key)]
    )
    monkeypatch.setattr(
        documents,
        "stat_object",
        lambda key, *, bucket: SimpleNamespace(size=len(objects[(bucket, key)])),
    )
    monkeypatch.setattr(
        documents, "delete_object", lambda key, *, bucket: objects.pop((bucket, key), None)
    )
    monkeypatch.setattr(
        indexer, "read_object", lambda key, *, bucket, max_size: objects[(bucket, key)]
    )
    monkeypatch.setattr(
        indexer, "embed_texts", lambda texts: ([[1.0, 0.0] for _ in texts], "test-embedding")
    )
    schedule = Mock()
    monkeypatch.setattr(documents, "_schedule_index_task", schedule)
    return {"base_id": base_id, "objects": objects, "headers": login(client), "schedule": schedule}


def create(client, setup, **overrides):
    body = {
        "knowledge_base_id": setup["base_id"],
        "title": "设备指南",
        "format": "md",
        "content": "# 设备指南\n\n请检查电源并启动设备。",
        "request_id": str(uuid4()),
        **overrides,
    }
    return client.post("/api/v1/documents/create", headers=setup["headers"], json=body), body


@pytest.mark.parametrize("format_name", ["md", "docx"])
def test_authorized_role_can_create_store_index_and_retry_once(app, client, creation, format_name):
    content = word_content() if format_name == "docx" else "# 设备指南\n\n请检查电源并启动设备。"
    response, body = create(client, creation, format=format_name, content=content)
    assert response.status_code == 201
    data = response.get_json()["data"]
    assert data["version"]["filename"] == f"设备指南.{format_name}"
    assert data["version"]["status"] == "UPLOADED"
    assert data["version"]["task"]["status"] == "PENDING"
    assert len(creation["objects"]) == 1
    payload = next(iter(creation["objects"].values()))
    if format_name == "md":
        assert payload.decode("utf-8") == content
    else:
        document = WordDocument(BytesIO(payload))
        assert document.paragraphs[0].style.name == "Heading 1"
        assert document.paragraphs[1].runs[0].bold and document.paragraphs[1].runs[0].underline
        assert document.tables[0].cell(1, 1).text == "绿色表示就绪"
        assert "绿色表示就绪" in _extract_text("guide.docx", payload)
    creation["schedule"].assert_called_once()
    replay = client.post("/api/v1/documents/create", headers=creation["headers"], json=body)
    assert replay.status_code == 200
    assert replay.get_json()["data"]["document"]["id"] == data["document"]["id"]
    with app.app_context():
        for model in [Document, DocumentVersion, IndexTask]:
            assert db.session.scalar(db.select(db.func.count()).select_from(model)) == 1
        assert (
            db.session.scalar(
                db.select(db.func.count())
                .select_from(AuditLog)
                .where(AuditLog.action == "document.create")
            )
            == 1
        )
    assert process_one(app)
    status = client.get(
        f"/api/v1/documents/created/{data['document']['id']}", headers=creation["headers"]
    )
    assert status.get_json()["data"]["version"]["status"] == "INDEXED"
    with app.app_context():
        document = db.session.get(Document, UUID(data["document"]["id"]))
        assert str(document.current_version_id) == data["version"]["id"]
        assert db.session.scalar(db.select(DocumentChunk)) is not None
    assert (
        client.get(
            f"/api/v1/documents/{data['document']['id']}", headers=creation["headers"]
        ).status_code
        == 403
    )


def test_creation_requires_grant_and_status_is_owner_scoped(app, client, creation):
    response, body = create(client, creation)
    doc_id = response.get_json()["data"]["document"]["id"]
    assert client.post("/api/v1/documents/create", json=body).status_code == 401
    result = app.test_cli_runner().invoke(
        args=["create-user", "--username", "other"],
        input="Writer-Password-2026!\nWriter-Password-2026!\n",
    )
    assert result.exit_code == 0
    other = login(client, "other")
    assert client.get(f"/api/v1/documents/created/{doc_id}", headers=other).status_code == 404
    assert client.post("/api/v1/documents/create", headers=other, json=body).status_code == 409
    with app.app_context():
        page = db.session.scalar(db.select(Page).where(Page.route_path == "/create-document"))
        page.status = "disabled"
        db.session.commit()
    assert create(client, creation)[0].status_code == 403


def test_failed_storage_rolls_back_and_draft_can_retry(app, client, creation, monkeypatch):
    original = documents.put_object
    monkeypatch.setattr(documents, "put_object", Mock(side_effect=StorageError("存储暂不可用")))
    response, body = create(client, creation)
    assert response.status_code == 503
    with app.app_context():
        assert db.session.get(Document, UUID(body["request_id"])) is None
        assert db.session.scalar(db.select(db.func.count()).select_from(IndexTask)) == 0
    monkeypatch.setattr(documents, "put_object", original)
    response = client.post("/api/v1/documents/create", headers=creation["headers"], json=body)
    assert response.status_code == 201
    assert len(creation["objects"]) == 1


def test_failed_completion_cleans_uploaded_object(app, client, creation, monkeypatch):
    monkeypatch.setattr(documents, "stat_object", lambda *args, **kwargs: SimpleNamespace(size=1))
    response, body = create(client, creation)
    assert response.status_code == 422
    assert not creation["objects"]
    with app.app_context():
        assert db.session.get(Document, UUID(body["request_id"])) is None


@pytest.mark.parametrize(
    "overrides",
    [
        {"content": ""},
        {"format": "pdf"},
        {"content": "\x00"},
        {"format": "docx", "content": {"type": "doc", "content": [{"type": []}]}},
        {
            "format": "docx",
            "content": {
                "type": "doc",
                "content": [
                    {
                        "type": "paragraph",
                        "content": [
                            {
                                "type": "text",
                                "text": "link",
                                "marks": [
                                    {"type": "link", "attrs": {"href": "javascript:alert(1)"}}
                                ],
                            }
                        ],
                    }
                ],
            },
        },
    ],
)
def test_creation_rejects_invalid_content_without_storage(client, creation, overrides):
    assert create(client, creation, **overrides)[0].status_code == 422
    assert not creation["objects"]


def test_disabled_knowledge_base_rejects_creation(app, client, creation):
    with app.app_context():
        base = db.session.get(KnowledgeBase, UUID(creation["base_id"]))
        base.status = "disabled"
        db.session.commit()
    assert create(client, creation)[0].status_code == 409
    assert not creation["objects"]


def test_documents_grant_enables_existing_upload_edit_download_and_archive(
    app, client, creation, monkeypatch
):
    with app.app_context():
        role = db.session.scalar(db.select(Role).where(Role.code == "user"))
        page = db.session.scalar(db.select(Page).where(Page.route_path == "/documents"))
        db.session.add(RolePagePermission(role_id=role.id, page_id=page.id, can_view=True))
        db.session.commit()
    monkeypatch.setattr(
        documents, "presign_put", lambda *args, **kwargs: "https://storage.test/upload"
    )
    monkeypatch.setattr(
        documents, "presign_get", lambda *args, **kwargs: "https://storage.test/download"
    )
    content = "# 上传文档\n\n正文".encode()
    headers = creation["headers"]
    response = client.post(
        "/api/v1/documents/presign-upload",
        headers=headers,
        json={
            "knowledge_base_id": creation["base_id"],
            "filename": "guide.md",
            "file_size": len(content),
        },
    )
    assert response.status_code == 201
    data = response.get_json()["data"]
    document_id, version_id = data["document"]["id"], data["version"]["id"]
    assert (
        client.put(
            f"/api/v1/documents/versions/{version_id}/upload",
            headers=headers,
            data=content,
            content_type="text/markdown",
        ).status_code
        == 202
    )
    assert process_one(app)
    assert (
        client.get(
            f"/api/v1/knowledge-bases/{creation['base_id']}/documents", headers=headers
        ).status_code
        == 200
    )
    assert client.get(f"/api/v1/documents/{document_id}", headers=headers).status_code == 200
    assert client.get(f"/api/v1/documents/{document_id}/content", headers=headers).get_json()[
        "data"
    ]["content"] == content.decode("utf-8")
    assert (
        client.put(
            f"/api/v1/documents/{document_id}", headers=headers, json={"title": "更新后的标题"}
        ).status_code
        == 200
    )
    assert (
        client.get(f"/api/v1/documents/{document_id}/download-url", headers=headers).status_code
        == 200
    )
    assert client.delete(f"/api/v1/documents/{document_id}", headers=headers).status_code == 202
