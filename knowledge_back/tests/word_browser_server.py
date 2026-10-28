"""Isolated browser-test API: real Flask routes/SQLite/DOCX, in-memory object storage.

Started only by word-editing.spec.ts; never imported by the application.
"""

import importlib
import json
import tempfile
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from conftest import app as app_fixture
from docx import Document
from flask import Response, request
from test_word_editing import external_word
from werkzeug.serving import make_server

with tempfile.TemporaryDirectory(prefix="kbs-word-browser-") as directory:
    monkeypatch = pytest.MonkeyPatch()
    app = app_fixture.__wrapped__(Path(directory), monkeypatch)
    app.config.update(INDEX_WORKER_AUTOSTART=False)
    documents = importlib.import_module("app.documents")
    indexer = importlib.import_module("app.indexer")
    objects = {}

    def put(key, stream, *, bucket, length, content_type):
        payload = stream.read()
        assert len(payload) == length
        objects[key] = payload

    def schedule():
        while indexer.process_one(app):
            pass

    monkeypatch.setattr(documents, "put_object", put)
    monkeypatch.setattr(documents, "read_object", lambda key, **kw: objects[key])
    monkeypatch.setattr(
        documents, "stat_object", lambda key, **kw: SimpleNamespace(size=len(objects[key]))
    )
    monkeypatch.setattr(documents, "delete_object", lambda key, **kw: objects.pop(key, None))
    monkeypatch.setattr(documents, "presign_put", lambda *a, **kw: "/fixture/requires-fallback")
    monkeypatch.setattr(documents, "_schedule_index_task", schedule)
    monkeypatch.setattr(indexer, "read_object", lambda key, **kw: objects[key])
    monkeypatch.setattr(
        indexer, "embed_texts", lambda texts: ([[1.0, 0.0] for _ in texts], "browser-test")
    )
    monkeypatch.setattr(
        importlib.import_module("app.retrieval"),
        "read_object",
        lambda key, **kw: objects[key],
        raising=False,
    )

    @app.get("/fixture/external.docx")
    def fixture():
        return Response(external_word(), content_type=documents.CONTENT_TYPES[".docx"])

    @app.get("/fixture/objects/<path:key>")
    def download(key):
        return Response(
            objects[key],
            headers={"Content-Disposition": 'attachment; filename="download.docx"'},
            content_type=documents.CONTENT_TYPES[".docx"],
        )

    @app.post("/fixture/inspect")
    def inspect_docx():
        doc = Document(BytesIO(request.get_data()))
        p = doc.paragraphs[1]
        return {
            "text": p.text,
            "bold": p.runs[0].bold,
            "italic": p.runs[0].italic,
            "color": str(p.runs[0].font.color.rgb),
            "tables": len(doc.tables),
            "page_width": doc.sections[0].page_width.cm,
            "header": doc.sections[0].header.paragraphs[0].text,
        }

    server = make_server("127.0.0.1", 0, app)
    origin = f"http://127.0.0.1:{server.server_port}"
    monkeypatch.setattr(
        documents, "presign_get", lambda key, *a, **kw: f"{origin}/fixture/objects/{key}"
    )
    print(json.dumps({"origin": origin}), flush=True)
    server.serve_forever()
