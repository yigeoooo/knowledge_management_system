"""Generate managed Markdown and Word files from editor content."""

from io import BytesIO
from urllib.parse import urlparse

from docx import Document as WordDocument
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from .http import ApiError
from .word_styles import (
    ALIGNMENTS,
    color,
    number,
    paragraph_style,
    points,
    text_style,
    validate_attributes,
    xml_value,
)

INLINE = {"text", "hardBreak", "datetime"}
BLOCKS = {
    "paragraph",
    "heading",
    "bulletList",
    "orderedList",
    "blockquote",
    "codeBlock",
    "horizontalRule",
    "table",
    "taskList",
    "pageBreak",
}
CHILDREN = {
    "doc": BLOCKS,
    "paragraph": INLINE,
    "heading": INLINE,
    "codeBlock": INLINE,
    "blockquote": BLOCKS,
    "bulletList": {"listItem"},
    "orderedList": {"listItem"},
    "listItem": BLOCKS,
    "taskList": {"taskItem"},
    "taskItem": BLOCKS,
    "table": {"tableRow"},
    "tableRow": {"tableCell", "tableHeader"},
    "tableCell": BLOCKS - {"table"},
    "tableHeader": BLOCKS - {"table"},
    "text": set(),
    "hardBreak": set(),
    "horizontalRule": set(),
    "pageBreak": set(),
    "datetime": set(),
}
MARKS = {
    "bold",
    "italic",
    "underline",
    "strike",
    "code",
    "link",
    "textStyle",
    "letterSpacing",
    "subscript",
    "superscript",
}
MAX_CHARACTERS = 200_000


def _invalid(message="Word 内容格式不正确"):
    raise ApiError(message, 422, 42202)


def _valid_text(text):
    if not isinstance(text, str) or any(
        (ord(char) < 32 and char not in "\n\r\t")
        or 0xD800 <= ord(char) <= 0xDFFF
        or ord(char) in {0xFFFE, 0xFFFF}
        for char in text
    ):
        _invalid("文档内容包含不支持的字符")


def _validate_word(root, *, allow_empty=False):
    nodes, characters, meaningful = 0, 0, False

    def visit(node, depth=0):
        nonlocal nodes, characters, meaningful
        nodes += 1
        if nodes > 10_000 or depth > 24:
            _invalid("文档结构过大或嵌套过深")
        if (
            not isinstance(node, dict)
            or not isinstance(node.get("type"), str)
            or node["type"] not in CHILDREN
        ):
            _invalid()
        kind = node["type"]
        attrs = node.get("attrs") or {}
        validate_attributes(kind, attrs)
        if kind == "heading" and (
            type(attrs.get("level", 1)) is not int or not 1 <= attrs.get("level", 1) <= 6
        ):
            _invalid()
        if kind == "orderedList" and (
            type(attrs.get("start", 1)) is not int or not 1 <= attrs.get("start", 1) <= 1_000_000
        ):
            _invalid()
        if kind in {"text", "datetime"}:
            text = node.get("text") if kind == "text" else attrs.get("text", "")
            _valid_text(text)
            characters += len(text)
            meaningful = meaningful or bool(text.strip())
        marks = node.get("marks", [])
        if not isinstance(marks, list):
            _invalid()
        for mark in marks:
            if (
                not isinstance(mark, dict)
                or not isinstance(mark.get("type"), str)
                or mark["type"] not in MARKS
            ):
                _invalid()
            validate_attributes(mark["type"], mark.get("attrs") or {})
            if mark["type"] == "link":
                attributes = mark.get("attrs") or {}
                if not isinstance(attributes, dict):
                    _invalid()
                href = attributes.get("href")
                _valid_text(href)
                try:
                    scheme = urlparse(href).scheme.lower()
                except ValueError:
                    _invalid("链接格式不正确")
                if scheme not in {"http", "https", "mailto"}:
                    _invalid("链接仅支持 http、https 或 mailto 地址")
        children = node.get("content", [])
        if not isinstance(children, list):
            _invalid()
        if kind in {"bulletList", "orderedList", "taskList", "table"} and not children:
            _invalid()
        for child in children:
            if (
                not isinstance(child, dict)
                or not isinstance(child.get("type"), str)
                or child["type"] not in CHILDREN[kind]
            ):
                _invalid()
            visit(child, depth + 1)
        if kind == "table":
            _table_layout(node)
        if kind == "doc":
            _page_settings(attrs.get("page"))
        if characters > MAX_CHARACTERS:
            _invalid(f"文档正文不能超过 {MAX_CHARACTERS} 个字符")

    if not isinstance(root, dict) or root.get("type") != "doc":
        _invalid()
    visit(root)
    if not meaningful and not allow_empty:
        _invalid("请先填写文档内容")


def _table_layout(node):
    rows = node.get("content", [])
    if not 1 <= len(rows) <= 200:
        _invalid("表格最多 200 行、20 列")
    occupied, cells = set(), []
    columns = 0
    for row_index, row in enumerate(rows):
        col = 0
        for cell in row.get("content", []):
            while (row_index, col) in occupied:
                col += 1
            attrs = cell.get("attrs") or {}
            width, height = attrs.get("colspan", 1), attrs.get("rowspan", 1)
            if col + width > 20 or row_index + height > len(rows):
                _invalid("表格合并范围越界，最多 200 行、20 列")
            for y in range(row_index, row_index + height):
                for x in range(col, col + width):
                    if (y, x) in occupied:
                        _invalid("表格合并单元格重叠")
                    occupied.add((y, x))
            cells.append((row_index, col, height, width, cell))
            columns = max(columns, col + width)
            col += width
    if not columns or len(occupied) != len(rows) * columns:
        _invalid("表格结构不完整，请修复表格后保存")
    return columns, cells


def _page_settings(settings, section=None):
    if settings is None:
        return
    if not isinstance(settings, dict):
        _invalid("页面设置不正确")
    size, margin = settings.get("size"), settings.get("margin")
    if not isinstance(size, dict) or not isinstance(margin, dict):
        _invalid("页面设置不正确")
    width, height = number(size.get("width"), 5, 100), number(size.get("height"), 5, 100)
    orientation = settings.get("orientation", "portrait")
    if orientation not in ("portrait", "landscape"):
        _invalid("页面方向不正确")
    if orientation == "landscape":
        width, height = height, width
    edges = {side: number(margin.get(side), 0, 20) for side in ("left", "right", "top", "bottom")}
    if edges["left"] + edges["right"] >= width or edges["top"] + edges["bottom"] >= height:
        _invalid("页边距不能超过纸张尺寸")
    if section is not None:
        section.orientation = (
            WD_ORIENT.LANDSCAPE if orientation == "landscape" else WD_ORIENT.PORTRAIT
        )
        section.page_width, section.page_height = Cm(width), Cm(height)
        for side, value in edges.items():
            setattr(section, f"{side}_margin", Cm(value))


def _inline(paragraph, nodes):
    for node in nodes:
        if node["type"] == "hardBreak":
            paragraph.add_run().add_break(WD_BREAK.LINE)
            continue
        run = paragraph.add_run(
            node["text"] if node["type"] == "text" else (node.get("attrs") or {}).get("text", "")
        )
        for mark in node.get("marks", []):
            kind = mark["type"]
            if kind in {"bold", "italic", "underline"}:
                setattr(run, kind, True)
            elif kind == "strike":
                run.font.strike = True
            elif kind == "code":
                run.font.name = "Consolas"
            elif kind == "textStyle":
                text_style(run, mark.get("attrs") or {})
            elif kind == "letterSpacing":
                spacing = (mark.get("attrs") or {}).get("spacing")
                if spacing is not None:
                    xml_value(
                        run._r.get_or_add_rPr(),
                        "spacing",
                        val=round(points(spacing, -10, 100) * 20),
                    )
            elif kind in {"subscript", "superscript"}:
                setattr(run.font, kind, True)
            elif kind == "link":
                from docx.opc.constants import RELATIONSHIP_TYPE

                relation = paragraph.part.relate_to(
                    mark["attrs"]["href"], RELATIONSHIP_TYPE.HYPERLINK, is_external=True
                )
                hyperlink = OxmlElement("w:hyperlink")
                hyperlink.set(qn("r:id"), relation)
                run.style = (
                    "Hyperlink"
                    if "Hyperlink" in paragraph.part.document.styles
                    else "Default Paragraph Font"
                )
                hyperlink.append(run._r)
                paragraph._p.append(hyperlink)


def _numbering(document, ordered, start, level, list_type=None):
    numbering = document.part.numbering_part.element
    existing = [
        int(item.get(qn("w:abstractNumId"))) for item in numbering.findall(qn("w:abstractNum"))
    ]
    abstract_id = max(existing, default=-1) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    for tag, value in [
        ("start", str(start)),
        (
            "numFmt",
            {
                "lower-alpha": "lowerLetter",
                "lower-latin": "lowerLetter",
                "upper-latin": "upperLetter",
                "decimal-leading-zero": "decimalZero",
                "trad-chinese-informal": "chineseCountingThousand",
                "simp-chinese-formal": "chineseLegalSimplified",
                "upper-alpha": "upperLetter",
                "lower-roman": "lowerRoman",
                "upper-roman": "upperRoman",
                "cjk-ideographic": "chineseCounting",
            }.get(list_type, "decimal")
            if ordered
            else "bullet",
        ),
        ("lvlText", "%1." if ordered else {"circle": "○", "square": "▪"}.get(list_type, "•")),
    ]:
        item = OxmlElement(f"w:{tag}")
        item.set(qn("w:val"), value)
        lvl.append(item)
    ppr = OxmlElement("w:pPr")
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), str(360 * (level + 1)))
    indent.set(qn("w:hanging"), "360")
    ppr.append(indent)
    lvl.append(ppr)
    abstract.append(lvl)
    numbering.append(abstract)
    return numbering.add_num(abstract_id).numId


def _word_bytes(content, source=None):
    _validate_word(content)
    from .word_import import load_word, save_snapshot

    document = load_word(source) if source else WordDocument()
    if source:
        # Retain styles, header/footer parts and relationships from the base
        # version. Only the editable body is replaced; the source stays immutable.
        for child in list(document.element.body):
            if child.tag != qn("w:sectPr"):
                document.element.body.remove(child)
    else:
        normal = document.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10.5)
        normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")

    _page_settings((content.get("attrs") or {}).get("page"), document.sections[0])

    def render(nodes, container, level=0):
        for node in nodes:
            kind, children = node["type"], node.get("content", [])
            attrs = node.get("attrs") or {}
            if kind in {"paragraph", "heading", "codeBlock"}:
                style = (
                    f"Heading {(node.get('attrs') or {}).get('level', 1)}"
                    if kind == "heading"
                    else "Normal"
                )
                paragraph = container.add_paragraph(style=style)
                _inline(paragraph, children)
                paragraph_style(paragraph, attrs)
                if kind == "codeBlock":
                    for run in paragraph.runs:
                        run.font.name = "Consolas"
                if level:
                    paragraph.paragraph_format.left_indent = Inches(0.25 * level)
            elif kind in {"bulletList", "orderedList", "taskList"}:
                num_id = _numbering(
                    document,
                    kind == "orderedList",
                    (node.get("attrs") or {}).get("start", 1),
                    level,
                    attrs.get("listType"),
                )
                for item in children:
                    item_blocks = item.get("content", [])
                    first = (
                        item_blocks[0]
                        if item_blocks and item_blocks[0]["type"] == "paragraph"
                        else None
                    )
                    paragraph = container.add_paragraph()
                    if kind == "taskList":
                        paragraph.add_run(
                            "☑ " if (item.get("attrs") or {}).get("checked") else "☐ "
                        )
                    if first:
                        paragraph_style(paragraph, first.get("attrs") or {})
                        _inline(paragraph, first.get("content", []))
                    if kind != "taskList":
                        numpr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
                        numpr.get_or_add_ilvl().val = 0
                        numpr.get_or_add_numId().val = num_id
                    render(item_blocks[1:] if first else item_blocks, container, level + 1)
            elif kind == "blockquote":
                render(children, container, level + 1)
            elif kind == "pageBreak":
                container.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            elif kind == "horizontalRule":
                paragraph = container.add_paragraph()
                borders = OxmlElement("w:pBdr")
                border = OxmlElement("w:bottom")
                border.set(qn("w:val"), "single")
                border.set(qn("w:sz"), "6")
                borders.append(border)
                paragraph._p.get_or_add_pPr().append(borders)
            elif kind == "table":
                columns, cells = _table_layout(node)
                table = container.add_table(rows=len(children), cols=columns)
                table.style = "Table Grid"
                # Merge on the empty grid, then populate each anchor once.
                for row, col, row_span, col_span, cell_node in cells:
                    if row_span > 1 or col_span > 1:
                        table.cell(row, col).merge(
                            table.cell(row + row_span - 1, col + col_span - 1)
                        )
                for row, col, _, col_span, cell_node in cells:
                    cell = table.cell(row, col)
                    cell_attrs = cell_node.get("attrs") or {}
                    empty = cell.paragraphs[0]._p
                    render(cell_node.get("content", []), cell)
                    if len(cell.paragraphs) > 1:
                        cell._tc.remove(empty)
                    if cell_attrs.get("background") and (fill := color(cell_attrs["background"])):
                        xml_value(cell._tc.get_or_add_tcPr(), "shd", val="clear", fill=fill)
                    align = (cell_attrs.get("align") or "").split("-")
                    if len(align) == 2:
                        cell.vertical_alignment = {
                            "top": WD_CELL_VERTICAL_ALIGNMENT.TOP,
                            "middle": WD_CELL_VERTICAL_ALIGNMENT.CENTER,
                            "bottom": WD_CELL_VERTICAL_ALIGNMENT.BOTTOM,
                        }[align[1]]
                    for paragraph in cell.paragraphs:
                        if paragraph.alignment is None and align[0] in ALIGNMENTS:
                            paragraph.alignment = ALIGNMENTS[align[0]]
                        for run in paragraph.runs:
                            if cell_node["type"] == "tableHeader":
                                run.bold = True
                            if (
                                cell_attrs.get("color")
                                and run.font.color.rgb is None
                                and (rgb := color(cell_attrs["color"]))
                            ):
                                run.font.color.rgb = RGBColor.from_string(rgb)
                    widths = cell_attrs.get("colwidth")
                    if widths and all(widths):
                        table.autofit = False
                        cell.width = Pt(sum(widths) * 0.75)
                        for offset in range(col_span):
                            table.columns[col + offset].width = Pt(widths[offset] * 0.75)

    render(content.get("content", []), document)
    save_snapshot(document, content)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def document_bytes(format_name, content, *, source=None):
    if format_name == "md":
        _valid_text(content)
        if not content.strip():
            _invalid("请先填写文档内容")
        if len(content) > MAX_CHARACTERS:
            _invalid(f"文档正文不能超过 {MAX_CHARACTERS} 个字符")
        return content.encode("utf-8")
    if format_name == "docx":
        return _word_bytes(content, source)
    if format_name == "xlsx":
        if source is not None:
            _invalid("Excel 在线编辑不需要传入源文件")
        from .excel_authoring import workbook_bytes

        return workbook_bytes(content)
    _invalid("文档格式只能为 Markdown、Word 或 Excel")
